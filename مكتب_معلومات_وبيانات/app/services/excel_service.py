# -*- coding: utf-8 -*-
"""
خدمة معالجة واستيراد وتصدير كشوفات Excel
Excel Processing & Deduplication Service
"""

import sys
import os
from datetime import datetime

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from database import get_db_connection
from services.passport_service import PassportService

try:
    import openpyxl
except ImportError:
    openpyxl = None


class ExcelService:
    @staticmethod
    def inspect_excel_file(file_path):
        """قراءة وفحص ملف Excel وكشف الأعمدة والتكرارات"""
        if not openpyxl:
            return False, "مكتبة openpyxl غير مثبتة. يرجى تثبيتها عبر pip install openpyxl"

        if not os.path.exists(file_path):
            return False, "ملف Excel غير موجود"

        wb = openpyxl.load_workbook(file_path, data_only=True)
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT passport_number FROM passports")
        existing_numbers = set(row[0].strip().lower() for row in cursor.fetchall())
        conn.close()

        results = {
            "total_rows": 0,
            "new_records": 0,
            "duplicate_records": 0,
            "rows": []
        }

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            header = [str(cell.value or '').strip().lower() for cell in sheet[1]]

            # مطابقة أسماء الأعمدة تلقائياً
            p_idx = -1
            name_idx = -1
            nat_idx = -1
            office_idx = -1

            for i, h in enumerate(header):
                if any(k in h for k in ['جواز', 'passport', 'رقم']):
                    p_idx = i
                elif any(k in h for k in ['اسم', 'name', 'صاحب']):
                    name_idx = i
                elif any(k in h for k in ['جنسية', 'nationality']):
                    nat_idx = i
                elif any(k in h for k in ['مكتب', 'وكالة', 'جهة', 'office']):
                    office_idx = i

            for row_cells in sheet.iter_rows(min_row=2, values_only=True):
                if not any(row_cells):
                    continue

                p_num = str(row_cells[p_idx] if p_idx >= 0 and p_idx < len(row_cells) and row_cells[p_idx] else '').strip()
                name = str(row_cells[name_idx] if name_idx >= 0 and name_idx < len(row_cells) and row_cells[name_idx] else 'بدون اسم').strip()
                nat = str(row_cells[nat_idx] if nat_idx >= 0 and nat_idx < len(row_cells) and row_cells[nat_idx] else 'يمني').strip()
                office = str(row_cells[office_idx] if office_idx >= 0 and office_idx < len(row_cells) and row_cells[office_idx] else '').strip()

                if p_num:
                    is_dup = p_num.lower() in existing_numbers
                    results["total_rows"] += 1
                    if is_dup:
                        results["duplicate_records"] += 1
                    else:
                        results["new_records"] += 1

                    results["rows"].append({
                        "passport_number": p_num,
                        "full_name": name,
                        "nationality": nat,
                        "office_name": office,
                        "is_duplicate": is_dup,
                        "sheet": sheet_name
                    })

        return True, results

    @staticmethod
    def bulk_import(rows, source_name="كشف Excel", officer="المشرف"):
        """استيراد السجلات إلى قاعدة البيانات مع توثيق الخط الزمني"""
        added = 0
        updated = 0

        for r in rows:
            p_num = r.get("passport_number")
            name = r.get("full_name")
            nat = r.get("nationality", "يمني")
            office = r.get("office_name", "")

            if not p_num:
                continue

            if r.get("is_duplicate"):
                PassportService.add_timeline_event(
                    passport_number=p_num,
                    event_title="ورود الجواز في كشف مكرر",
                    details=f"تكرر الجواز في {source_name} - المكتب: {office}",
                    officer=officer
                )
                updated += 1
            else:
                success, _ = PassportService.register_passport(
                    passport_number=p_num,
                    full_name=name,
                    nationality=nat,
                    office_name=office,
                    notes=f"مستورد من {source_name}",
                    officer=officer
                )
                if success:
                    added += 1

        return added, updated
